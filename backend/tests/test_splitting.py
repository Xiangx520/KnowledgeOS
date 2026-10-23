"""Structural boundaries, lossless code bodies, and bounded ingestion output."""
from collections import Counter

import pytest
from app.config import get_settings
from app.services.document_structure import ParsedBlock, ParsedPage, markdown_blocks
from app.services.evidence import clean_text
from app.services.parsing import chunk_pages, parse_document
from app.services.splitting import CodeBlockSplitter, HeadingSplitter, TableSplitter


@pytest.fixture
def small_chunks(monkeypatch):
    monkeypatch.setenv("RAG_CHUNK_SIZE", "100")
    monkeypatch.setenv("RAG_CHUNK_OVERLAP", "10")
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def split_markdown(text, page=7):
    return chunk_pages([ParsedPage(text, page, markdown_blocks(text))])


def code_body(chunk):
    return chunk[chunk.index("\n") + 1:chunk.rindex("\n")]


def assert_chunks(chunks, size=100):
    assert all(0 < len(c["text"]) <= size for c in chunks)
    assert [c["ordinal"] for c in chunks] == list(range(len(chunks)))
    assert all(set(c) == {"text", "page", "ordinal"} for c in chunks)


def test_heading_sections_and_fenced_false_positives(small_chunks):
    text = ("前言\n\n# 一级\n正文\n## 子级\n子正文\n子标题\n---\n更多正文\n"
            "~~~python\n# 不是标题\n| a | b |\n| --- | --- |\n~~~\n# 下节\n结束")
    blocks = markdown_blocks(text)
    sections = HeadingSplitter().split(blocks)
    assert [s[0].kind for s in sections] == ["text", "heading", "heading", "heading", "heading"]
    assert [b.heading_level for b in blocks if b.kind == "heading"] == [1, 2, 2, 1]
    assert not any(b.kind == "table" for b in blocks)
    chunks = split_markdown(text)
    assert_chunks(chunks)
    assert all(c["page"] == 7 for c in chunks)
    assert not any("前言" in c["text"] and "一级" in c["text"] for c in chunks)
    assert not any("子正文" in c["text"] and "子标题" in c["text"] for c in chunks)
    assert any("# 不是标题" in c["text"] and c["text"].startswith("```python") for c in chunks)


@pytest.mark.parametrize("fence", ["```", "~~~~", "````"])
@pytest.mark.parametrize("closed", [True, False])
def test_code_preserves_links_indentation_blank_lines_and_fences(small_chunks, fence, closed):
    body = "    [literal](file.md)\n\n    ![image](x.png)\n" + "    value = 你好世界\n" * 16 + "  tail  "
    raw = fence + "python\n" + body + ("\n" + fence if closed else "")
    cleaned = clean_text(raw)
    assert "[literal](file.md)" in cleaned and "![image](x.png)" in cleaned
    blocks = markdown_blocks(cleaned)
    assert len(blocks) == 1 and blocks[0].kind == "code"
    values = CodeBlockSplitter(100).split(blocks[0])
    assert all(len(v) <= 100 and v.startswith("```python\n") and v.endswith("\n```") for v in values)
    recovered = "".join(code_body(v) for v in values)
    # A renderer may insert one final newline before a closing fence.
    assert recovered.rstrip("\n") == blocks[0].text.rstrip("\n")


def test_code_embedded_fences_and_oversized_line(small_chunks):
    block = ParsedBlock("    " + "中文" * 140 + "\n```\nend\n", "code", language="text")
    chunks = CodeBlockSplitter(100).split(block)
    assert all(len(c) <= 100 for c in chunks)
    assert "".join(code_body(c) for c in chunks) == block.text


def test_table_rows_repeat_only_header_and_keep_surrounding_prose(small_chunks):
    rows = [f"| 编号{i:02} | 值{i:02} |" for i in range(25)]
    text = "表前\n\n| 编号 | 值 |\n| --- | --- |\n" + "\n".join(rows) + "\n\n表后"
    chunks = split_markdown(text)
    assert_chunks(chunks)
    assert chunks[0]["text"] == "表前" and chunks[-1]["text"] == "表后"
    tables = [c["text"] for c in chunks[1:-1]]
    assert len(tables) > 1
    assert all(t.startswith("| 编号 | 值 |\n| --- | --- |\n") for t in tables)
    counts = Counter(line for t in tables for line in t.splitlines()[2:])
    assert counts == Counter(rows)


def test_table_long_row_and_header_fallback(small_chunks):
    header = ("a | b", "--- | ---")
    row = "超长数据" * 100
    block = ParsedBlock("\n".join((*header, row)), "table", header=header)
    values = TableSplitter(100).split(block)
    assert all(len(v) <= 100 for v in values)
    assert "".join(v[len("\n".join(header)) + 1:] for v in values) == row
    giant_header = ("标题" * 100, "--- | ---")
    giant = ParsedBlock("\n".join((*giant_header, row)), "table", header=giant_header)
    assert "".join(TableSplitter(100).split(giant)) == giant.text
    giant_code = ParsedBlock(row, "code", language="x" * 100)
    from app.services.document_structure import code_text
    assert "".join(CodeBlockSplitter(100).split(giant_code)) == code_text(giant_code)


def test_plain_text_pages_empty_nul_and_determinism(small_chunks):
    pages = [ParsedPage("普通文字" * 60 + "\x00", 3), ParsedPage("末页", 9)]
    chunks = chunk_pages(pages)
    assert_chunks(chunks)
    assert chunks == chunk_pages(pages)
    assert chunks[0]["page"] == 3 and chunks[-1]["page"] == 9
    assert not any("\x00" in c["text"] for c in chunks)
    for empty in ([], [ParsedPage(" \x00 ")]):
        with pytest.raises(ValueError, match="没有可检索文本"):
            chunk_pages(empty)


def test_html_structure_source_order_and_no_duplicate_code(tmp_path, small_chunks):
    path = tmp_path / "ordered.html"
    path.write_text('<head><title>忽略</title></head><script>忽略</script><h1>标题</h1>'
                    '<p>表前 <b>加粗</b></p><table><tr><th>字段</th><th>数值</th></tr>'
                    '<tr><td>额度</td><td>500</td></tr></table><p>表后</p>'
                    '<pre><code class="language-python">  x = "[a](b)"\n\n  y = 2\n</code></pre>', encoding="utf-8")
    page = parse_document(path, path.name)[0]
    assert [b.kind for b in page.blocks] == ["heading", "text", "table", "text", "code"]
    assert page.text.index("表前") < page.text.index("额度") < page.text.index("表后")
    assert "忽略" not in page.text
    chunks = chunk_pages([page])
    assert_chunks(chunks)
    assert sum('x = "[a](b)"' in c["text"] for c in chunks) == 1
    assert chunks[-1]["text"].startswith("```python\n  x")


def test_docx_source_order_heading_and_explicit_code(tmp_path, small_chunks):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    doc = Document()
    doc.add_heading("标题", 1)
    doc.add_paragraph("表前")
    table = doc.add_table(rows=2, cols=2)
    for cell, value in zip([c for row in table.rows for c in row.cells], ["字段", "数值", "额度", "500"]):
        cell.text = value
    doc.add_paragraph("表后")
    doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("  x = 1", "Code Block")
    doc.add_paragraph("", "Code Block")
    doc.add_paragraph("    y = 2", "Code Block")
    doc.add_paragraph("普通段落 x = 3")
    path = tmp_path / "ordered.docx"
    doc.save(path)
    page = parse_document(path, path.name)[0]
    assert [b.kind for b in page.blocks] == ["heading", "text", "table", "text", "code", "text"]
    assert page.blocks[0].heading_level == 1
    assert page.blocks[4].text == "  x = 1\n\n    y = 2"
    assert page.text.index("表前") < page.text.index("额度") < page.text.index("表后")
    chunks = chunk_pages([page])
    assert_chunks(chunks)
    assert all(c["page"] == 1 for c in chunks)


def test_markdown_file_cleans_only_prose(tmp_path, small_chunks):
    path = tmp_path / "mixed.md"
    path.write_text("* [目录](other.md)\n正文[链接](https://example.org)\n\n"
                    "| 字段 | 值 |\n| --- | --- |\n| 链接 | [原文](other.md) |\n\n"
                    "```\n    [保留](other.md)\n    ![保留](image.png)\n```", encoding="utf-8")
    page = parse_document(path, path.name)[0]
    assert "目录" not in page.text and "正文链接" in page.text
    assert "[原文](other.md)" in page.blocks[1].text
    assert page.blocks[-1].text == "    [保留](other.md)\n    ![保留](image.png)\n"
    assert_chunks(chunk_pages([page]))
